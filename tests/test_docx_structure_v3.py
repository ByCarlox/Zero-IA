"""DOCX semantic metadata must preserve the canonical text and UTF-16 ranges."""
import io
import unittest
import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.document_parser import extract_from_docx, parse_document


def as_bytes(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDocxStructureV3(unittest.TestCase):
    def test_utf16_and_semantic_order(self):
        document = docx.Document()
        document.add_heading('Introducción 🧪', level=1)
        document.add_heading('Contexto', level=2)
        document.add_paragraph('Texto humano y verificable.')
        document.add_paragraph('Primer punto', style='List Bullet')
        document.add_paragraph('Introducción\t3')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Valor'
        table.cell(0, 1).text = '42'
        document.add_paragraph('Después de la tabla.')
        result = parse_document(as_bytes(document), 'ejemplo.docx')
        blocks = result['extraction']['structure']['blocks']
        self.assertEqual([b['kind'] for b in blocks], ['heading', 'heading', 'body', 'list', 'toc', 'table', 'body'])
        self.assertEqual([blocks[0]['level'], blocks[1]['level']], [1, 2])
        self.assertEqual(result['full_text'], '\n\n'.join(result['paragraphs']))
        encoded = result['full_text'].encode('utf-16-le')
        for block in blocks:
            self.assertEqual(encoded[block['start'] * 2:block['end'] * 2].decode('utf-16-le'), block['text'])
        self.assertEqual(blocks[1]['start'], len(blocks[0]['text']) + 3)

    def test_outline_levels_and_inheritance(self):
        document = docx.Document()
        style = document.styles.add_style('Custom section', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles['Heading 2']
        document.add_paragraph('Inherited heading', style=style)
        paragraph = document.add_paragraph('Explicit outline')
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '2')
        paragraph._p.get_or_add_pPr().append(outline)
        plain = document.add_paragraph('Override heading', style='Heading 1')
        override = OxmlElement('w:outlineLvl')
        override.set(qn('w:val'), '9')
        plain._p.get_or_add_pPr().append(override)
        result = extract_from_docx(as_bytes(document))
        self.assertEqual([b.get('level') for b in result['structure']['blocks']], [2, 3, None])
        self.assertEqual(result['structure']['blocks'][2]['kind'], 'body')

    def test_numbering_and_toc_style(self):
        document = docx.Document()
        paragraph = document.add_paragraph('Numbered item')
        num = OxmlElement('w:numPr')
        paragraph._p.get_or_add_pPr().append(num)
        style = document.styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph('Index entry without page', style=style)
        self.assertEqual([b['kind'] for b in extract_from_docx(as_bytes(document))['structure']['blocks']], ['list', 'toc'])


if __name__ == '__main__':
    unittest.main()
