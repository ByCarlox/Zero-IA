"""
Tests unitarios para la segmentación de oraciones y el parser de documentos.
"""

import unittest
import io
import docx
from core.sentence_tokenizer import split_into_sentences, split_into_paragraphs
from core.document_parser import extract_from_docx, extract_from_txt


class TestTokenizerAndParser(unittest.TestCase):
    def test_sentence_splitting_with_abbreviations(self):
        text = "El Dr. Gómez analizó el art. 5 en la pág. 12. Luego, concluyó el experimento con éxito."
        sentences = split_into_sentences(text)
        self.assertEqual(len(sentences), 2)
        self.assertIn("Dr. Gómez", sentences[0])
        self.assertIn("pág. 12.", sentences[0])
        self.assertEqual(sentences[1], "Luego, concluyó el experimento con éxito.")

    def test_split_into_paragraphs(self):
        raw = "Párrafo uno.\n\nPárrafo dos con más texto.\n\nPárrafo tres."
        paragraphs = split_into_paragraphs(raw)
        self.assertEqual(len(paragraphs), 3)

    def test_extract_from_docx(self):
        doc = docx.Document()
        doc.add_paragraph("Introducción al proyecto de máster.")
        doc.add_paragraph("Segunda línea con metodología analítica.")
        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        result = extract_from_docx(docx_bytes)
        self.assertEqual(result["format"], "docx")
        self.assertEqual(len(result["paragraphs"]), 2)
        self.assertIn("Introducción al proyecto", result["full_text"])

    def test_extract_from_txt(self):
        raw_bytes = "Este es un texto simple de prueba.\n\nSegundo bloque.".encode("utf-8")
        result = extract_from_txt(raw_bytes)
        self.assertEqual(result["format"], "txt")
        self.assertEqual(len(result["paragraphs"]), 2)


if __name__ == "__main__":
    unittest.main()
