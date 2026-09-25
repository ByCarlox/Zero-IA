"""Exports from the canonical report; no independent scoring or text rewriting."""
import io
import docx
from docx.enum.text import WD_COLOR_INDEX
from core.engine_bridge import call_engine


def export_markdown_report(analysis_result):
    return call_engine('report', analysis=analysis_result)


def export_annotated_docx(analysis_result):
    doc = docx.Document()
    doc.add_heading('Validador Académico · Zero-IA', 0)
    doc.add_paragraph('Exportación del texto extraído. No reproduce la maquetación del archivo original.')
    if analysis_result.get('validationScore') is not None:
        doc.add_paragraph(f"Índice de validación: {analysis_result['validationScore']}% ({analysis_result.get('cleanPercentage', 0)}% texto libre de incidencias).")
    doc.add_heading('Texto anotado', 1)
    source = analysis_result['sourceText'].encode('utf-16-le')
    p = doc.add_paragraph()
    cursor = 0
    for sentence in analysis_result['sentences']:
        start, end = sentence['start'] * 2, sentence['end'] * 2
        p.add_run(source[cursor:start].decode('utf-16-le'))
        run = p.add_run(source[start:end].decode('utf-16-le'))
        if sentence['riskLevel'] == 'high':
            run.font.highlight_color = WD_COLOR_INDEX.PINK
        elif sentence['riskLevel'] == 'medium':
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        cursor = end
    p.add_run(source[cursor:].decode('utf-16-le'))
    doc.add_heading('Informe', 1)
    for line in export_markdown_report(analysis_result).splitlines():
        doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
